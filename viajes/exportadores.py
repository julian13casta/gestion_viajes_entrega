"""
Funciones de exportación de datos filtrados a CSV y a Excel (XLSX).

Se mantienen separadas de `views.py` para que las vistas solo se ocupen de
orquestar request → queryset → respuesta HTTP, y esta capa se ocupe
exclusivamente de "cómo se ve" cada formato de exportación.
"""

import csv
from datetime import datetime

from django.http import HttpResponse

COLUMNAS = [
    ("codigo", "Código"),
    ("placa", "Placa"),
    ("tipo_vehiculo", "Tipo de vehículo"),
    ("cliente", "Cliente"),
    ("origen", "Origen"),
    ("destino", "Destino"),
    ("fecha_inicio", "Fecha inicio"),
    ("fecha_fin", "Fecha fin"),
    ("numero_entregas", "N° entregas"),
    ("facturacion", "Facturación"),
    ("validado", "Validado"),
    ("observacion", "Observación"),
]


def _nombre_archivo(extension: str) -> str:
    marca = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"viajes_export_{marca}.{extension}"


def _fila_de(vehiculo) -> list:
    return [
        vehiculo.codigo,
        vehiculo.placa,
        vehiculo.tipo_vehiculo,
        vehiculo.cliente,
        vehiculo.origen,
        vehiculo.destino,
        vehiculo.fecha_inicio.strftime("%Y-%m-%d"),
        vehiculo.fecha_fin.strftime("%Y-%m-%d") if vehiculo.fecha_fin else "",
        vehiculo.numero_entregas,
        float(vehiculo.facturacion),
        "Sí" if vehiculo.validado else "No",
        vehiculo.observacion,
    ]


def exportar_csv(queryset) -> HttpResponse:
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{_nombre_archivo("csv")}"'
    # BOM para que Excel en Windows detecte UTF-8 y no rompa tildes/ñ.
    response.write("﻿")

    writer = csv.writer(response)
    writer.writerow([etiqueta for _campo, etiqueta in COLUMNAS])
    for vehiculo in queryset.iterator():
        writer.writerow(_fila_de(vehiculo))
    return response


def exportar_excel(queryset) -> HttpResponse:
    # Import local: openpyxl solo es necesario aquí, así el resto de la app
    # no depende de una librería pesada si algún día se retira este export.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Viajes"

    encabezado_relleno = PatternFill(start_color="F78C2D", end_color="F78C2D", fill_type="solid")
    encabezado_fuente = Font(color="FFFFFF", bold=True)

    for col_idx, (_campo, etiqueta) in enumerate(COLUMNAS, start=1):
        celda = ws.cell(row=1, column=col_idx, value=etiqueta)
        celda.fill = encabezado_relleno
        celda.font = encabezado_fuente
        celda.alignment = Alignment(horizontal="center")

    fila_idx = 2
    verde = PatternFill(start_color="E3F5E8", end_color="E3F5E8", fill_type="solid")
    rojo = PatternFill(start_color="FCE8E6", end_color="FCE8E6", fill_type="solid")

    for vehiculo in queryset.iterator():
        fila = _fila_de(vehiculo)
        for col_idx, valor in enumerate(fila, start=1):
            ws.cell(row=fila_idx, column=col_idx, value=valor)
        # Color de fila según estado de validación (coherente con la tabla web).
        relleno = verde if vehiculo.validado else rojo
        for col_idx in range(1, len(COLUMNAS) + 1):
            ws.cell(row=fila_idx, column=col_idx).fill = relleno
        fila_idx += 1

    anchos = [12, 10, 16, 26, 14, 14, 13, 13, 11, 14, 10, 40]
    for col_idx, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = ancho

    ws.freeze_panes = "A2"

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{_nombre_archivo("xlsx")}"'
    wb.save(response)
    return response
